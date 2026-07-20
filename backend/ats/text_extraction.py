"""
Text extraction helpers for resume files (PDF, DOCX, plain text).
"""
import os


def extract_resume_text(file_path):
    """
    Extract raw text from a resume file based on its extension.
    Returns an empty string if extraction fails for any reason, rather
    than raising, so the calling pipeline can handle the failure cleanly.
    """
    ext = os.path.splitext(file_path)[1].lower()

    try:
        if ext == ".pdf":
            return _extract_pdf(file_path)
        elif ext == ".docx":
            return _extract_docx(file_path)
        elif ext == ".doc":
            # Legacy .doc binary format isn't supported by python-docx.
            # We still attempt a best-effort raw read so the pipeline
            # doesn't crash; quality will be poor for genuine .doc files.
            return _extract_fallback_text(file_path)
        else:
            return _extract_fallback_text(file_path)
    except Exception:
        return ""


def _extract_pdf(file_path):
    from pdfminer.high_level import extract_text
    text = extract_text(file_path) or ""
    return text


def _extract_docx(file_path):
    import docx
    doc = docx.Document(file_path)
    parts = [p.text for p in doc.paragraphs]

    # Tables often hold contact info / skill matrices in resumes.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts)


def _extract_fallback_text(file_path):
    with open(file_path, "r", errors="ignore") as f:
        return f.read()
